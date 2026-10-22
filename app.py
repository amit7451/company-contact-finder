from __future__ import annotations

import csv
import io
import json
import os
import re
import threading
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from typing import Any
from urllib.parse import urlparse

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from scraper import normalize_input_url, scrape_website

APP_DIR = os.path.dirname(os.path.abspath(__file__))
MAX_ROWS = 1000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024

app = FastAPI(title="Company Contact Finder", version="1.0.0")
app.mount("/static", StaticFiles(directory=os.path.join(APP_DIR, "static")), name="static")

jobs: dict[str, dict[str, Any]] = {}
jobs_lock = threading.Lock()
job_pool = ThreadPoolExecutor(max_workers=2, thread_name_prefix="job")


def _looks_like_url(value: str) -> bool:
    value = (value or "").strip()
    if not value:
        return False
    if value.startswith(("http://", "https://", "www.")):
        return True
    return bool(re.fullmatch(r"[A-Za-z0-9.-]+\.[A-Za-z]{2,}(?:/[^\s]*)?", value))


def _normal_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _parse_sheet(data: bytes, filename: str) -> list[dict[str, str]]:
    ext = os.path.splitext(filename.lower())[1]
    if ext in (".xlsx", ".xlsm"):
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
        workbook.close()
        if not raw_rows:
            return []
        rows = [["" if v is None else str(v).strip() for v in row] for row in raw_rows]
    elif ext in (".csv", ".tsv", ".txt"):
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
        rows = [[str(v).strip() for v in row] for row in rows]
    else:
        raise HTTPException(status_code=400, detail="Upload a .csv, .tsv, .txt, .xlsx, or .xlsm file.")

    rows = [r for r in rows if any(str(v).strip() for v in r)]
    if not rows:
        return []
    header_map = [_normal_header(v) for v in rows[0]]
    company_keys = {"company", "companies", "companyname", "business", "businessname", "brand", "brandname", "organization", "organisation", "name", "client"}
    website_keys = {"website", "websiteurl", "url", "domain", "site", "companywebsite", "webaddress", "homepage"}
    company_idx = next((i for i, h in enumerate(header_map) if h in company_keys), None)
    website_idx = next((i for i, h in enumerate(header_map) if h in website_keys), None)
    # If the first row looks like headers, use it; otherwise treat all rows as data.
    header_likely = any(h in company_keys | website_keys | {"email", "phone", "notes", "industry"} for h in header_map)
    data_rows = rows[1:] if header_likely else rows
    if company_idx is None and website_idx is None:
        # Single-column lists are commonly just company names or domains.
        company_idx = 0
    output: list[dict[str, str]] = []
    for row in data_rows:
        def cell(idx: int | None) -> str:
            return str(row[idx]).strip() if idx is not None and idx < len(row) and row[idx] is not None else ""
        company = cell(company_idx)
        website = cell(website_idx)
        if not company and website:
            company = website
        if not company and not website:
            continue
        output.append({"company": company, "website": website})
        if len(output) >= MAX_ROWS:
            break
    return output


def _search_company_website(company: str) -> tuple[str | None, str]:
    """Best-effort name -> official site lookup using the DDGS search library."""
    # Optional official API support is useful if the public search provider throttles.
    brave_key = os.getenv("BRAVE_SEARCH_API_KEY", "").strip()
    if brave_key:
        try:
            import requests
            resp = requests.get(
                "https://api.search.brave.com/res/v1/web/search",
                params={"q": f'"{company}" official website', "count": 8},
                headers={"Accept": "application/json", "X-Subscription-Token": brave_key},
                timeout=8,
            )
            resp.raise_for_status()
            results = resp.json().get("web", {}).get("results", [])
            candidates = [(r.get("url", ""), r.get("title", ""), r.get("description", "")) for r in results]
        except Exception:
            candidates = []
    else:
        candidates = []
    if not candidates:
        try:
            from ddgs import DDGS
            with DDGS() as ddgs:
                results = ddgs.text(f'"{company}" official website', region="in-en", max_results=8)
            candidates = [(r.get("href") or r.get("url") or "", r.get("title", ""), r.get("body", "")) for r in (results or [])]
        except Exception as exc:
            return None, f"Company-name search failed ({exc.__class__.__name__}); provide the website URL or configure search."

    company_tokens = [t for t in re.findall(r"[a-z0-9]+", company.lower()) if len(t) > 2 and t not in {"the", "and", "for", "inc", "llc", "ltd", "limited", "private", "pvt"}]
    blocked = ("facebook.com", "instagram.com", "linkedin.com", "youtube.com", "wikipedia.org", "justdial.com", "indiamart.com", "sulekha.com", "crunchbase.com", "glassdoor.com", "ambitionbox.com", "twitter.com", "x.com", "reddit.com")
    best_url, best_score = None, -999
    for raw_url, title, snippet in candidates:
        url = normalize_input_url(raw_url)
        if not url:
            continue
        host = (urlparse(url).hostname or "").lower().removeprefix("www.")
        if any(b in host for b in blocked):
            continue
        haystack = f"{host} {title} {snippet}".lower()
        score = sum(3 if token in host else 1 for token in company_tokens if token in haystack)
        if "official" in haystack:
            score += 1
        if any(term in host for term in ("directory", "directory", "review", "jobs", "careers")):
            score -= 4
        if score > best_score:
            best_score, best_url = score, url
    if best_url and best_score >= 1:
        return best_url, "Website found by best-effort company-name search; please verify the domain before outreach."
    return None, "No likely official website found. Enter a website URL to improve the match."


def _process_one(item: dict[str, str]) -> dict[str, Any]:
    company = (item.get("company") or "").strip()
    website_input = (item.get("website") or "").strip()
    search_note = ""
    if not website_input and _looks_like_url(company):
        website_input = company
    if not website_input:
        website_input, search_note = _search_company_website(company)
    normalized = normalize_input_url(website_input or "")
    if not normalized:
        return {
            "company": company or website_input,
            "website": website_input,
            "site_title": "",
            "emails": [],
            "phones": [],
            "pages_visited": 0,
            "source_urls": [],
            "status": "website_not_found" if not website_input else "invalid_url",
            "error": search_note or "Could not identify a valid website. Add the official URL and retry.",
        }
    result = scrape_website(normalized)
    result["company"] = company or result.get("site_title") or normalized
    result["website"] = normalized
    if search_note:
        result["search_note"] = search_note
    return result


def _run_job(job_id: str, rows: list[dict[str, str]]) -> None:
    with jobs_lock:
        jobs[job_id]["status"] = "running"
        jobs[job_id]["message"] = "Scanning public pages and extracting listed business contacts."
    results: list[dict[str, Any] | None] = [None] * len(rows)
    try:
        with ThreadPoolExecutor(max_workers=4, thread_name_prefix="site") as pool:
            future_map = {pool.submit(_process_one, row): i for i, row in enumerate(rows)}
            for future in as_completed(future_map):
                index = future_map[future]
                try:
                    result = future.result()
                except Exception as exc:
                    result = {
                        "company": rows[index].get("company", ""),
                        "website": rows[index].get("website", ""),
                        "site_title": "", "emails": [], "phones": [], "pages_visited": 0,
                        "source_urls": [], "status": "error", "error": f"Unexpected error: {exc.__class__.__name__}",
                    }
                results[index] = result
                with jobs_lock:
                    jobs[job_id]["completed"] += 1
                    jobs[job_id]["results"] = [r for r in results if r is not None]
        with jobs_lock:
            jobs[job_id]["results"] = [r for r in results if r is not None]
            jobs[job_id]["status"] = "completed"
            jobs[job_id]["message"] = "Scan finished. Review the source URLs before contacting anyone."
            jobs[job_id]["finished_at"] = datetime.now(timezone.utc).isoformat()
    except Exception as exc:
        with jobs_lock:
            jobs[job_id]["status"] = "failed"
            jobs[job_id]["message"] = f"Job failed ({exc.__class__.__name__})."


def _job_snapshot(job: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": job["id"], "status": job["status"], "total": job["total"],
        "completed": job["completed"], "message": job["message"],
        "results": job.get("results", []),
    }


@app.get("/")
def index():
    return FileResponse(os.path.join(APP_DIR, "static", "index.html"))


@app.get("/api/health")
def health():
    return {"ok": True, "app": "Company Contact Finder"}


@app.post("/api/jobs")
async def create_job(items_json: str = Form(default="[]"), file: UploadFile | None = File(default=None)):
    rows: list[dict[str, str]] = []
    if items_json and items_json != "[]":
        try:
            raw = json.loads(items_json)
            if isinstance(raw, str):
                raw = raw.splitlines()
            for entry in raw:
                if isinstance(entry, str):
                    value = entry.strip()
                    if value:
                        rows.append({"company": value, "website": ""})
                elif isinstance(entry, dict):
                    company = str(entry.get("company") or entry.get("name") or "").strip()
                    website = str(entry.get("website") or entry.get("url") or entry.get("domain") or "").strip()
                    if company or website:
                        rows.append({"company": company or website, "website": website})
        except Exception:
            raise HTTPException(status_code=400, detail="Could not parse the input list.")
    if file and file.filename:
        data = await file.read(MAX_UPLOAD_BYTES + 1)
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=413, detail="File is larger than 10 MB.")
        rows.extend(_parse_sheet(data, file.filename))
    # Stable de-duplication avoids scanning repeated rows in a batch.
    deduped: list[dict[str, str]] = []
    seen = set()
    for row in rows:
        key = ((row.get("company") or "").strip().lower(), (row.get("website") or "").strip().lower())
        if key not in seen and any(key):
            deduped.append(row)
            seen.add(key)
    rows = deduped
    if not rows:
        raise HTTPException(status_code=400, detail="Enter at least one company/website or upload a populated file.")
    if len(rows) > MAX_ROWS:
        rows = rows[:MAX_ROWS]
    job_id = uuid.uuid4().hex[:12]
    job = {
        "id": job_id, "status": "queued", "total": len(rows), "completed": 0,
        "message": "Queued for scanning.", "results": [], "created_at": datetime.now(timezone.utc).isoformat(),
    }
    with jobs_lock:
        jobs[job_id] = job
    job_pool.submit(_run_job, job_id, rows)
    return {"id": job_id, "total": len(rows), "status": "queued"}


@app.get("/api/jobs/{job_id}")
def get_job(job_id: str):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found. It may have expired when the app restarted.")
        return _job_snapshot(job)


def _flatten(result: dict[str, Any]) -> dict[str, Any]:
    return {
        "Company": result.get("company", ""),
        "Website": result.get("website", ""),
        "Page title": result.get("site_title", ""),
        "Emails": "; ".join(result.get("emails", []) or []),
        "Phone numbers": "; ".join(result.get("phones", []) or []),
        "Pages visited": result.get("pages_visited", 0),
        "Status": result.get("status", ""),
        "Notes": result.get("error", "") or result.get("search_note", ""),
        "Source URLs": "; ".join(result.get("source_urls", []) or []),
    }


@app.get("/api/jobs/{job_id}/download")
def download_job(job_id: str, format: str = "csv"):
    with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found.")
        results = list(job.get("results", []))
    if not results:
        raise HTTPException(status_code=409, detail="No results are available yet.")
    flat_rows = [_flatten(r) for r in results]
    headers = list(flat_rows[0].keys())
    if format.lower() == "xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Contact Results"
        sheet.append(headers)
        for item in flat_rows:
            sheet.append([item.get(h, "") for h in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="193B57")
            cell.alignment = Alignment(vertical="top")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        widths = [24, 36, 30, 38, 28, 14, 20, 45, 65]
        for i, width in enumerate(widths, start=1):
            sheet.column_dimensions[chr(64 + i)].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        buffer = io.BytesIO()
        workbook.save(buffer)
        buffer.seek(0)
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="contact-results-{job_id}.xlsx"'},
        )
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=headers)
    writer.writeheader()
    writer.writerows(flat_rows)
    return StreamingResponse(
        iter(["\ufeff" + buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="contact-results-{job_id}.csv"'},
    )
